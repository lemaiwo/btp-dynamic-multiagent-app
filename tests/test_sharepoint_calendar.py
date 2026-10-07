"""Calendar views of ``agents/sharepoint_workbook.py``.

Covered: runs, the code map and its default-deny, unmapped codes, skipped
rows, the stop marker, the lookup, conflicts, a window over New Year, and the
refusals (no cached values, dates not found, no sheet for a year).

Run:  python -m pytest tests/test_sharepoint_calendar.py
"""

from __future__ import annotations

import copy
import json
import sys
from datetime import date, datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

import pytest  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

from agents import sharepoint_workbook  # noqa: E402
from agents.sharepoint_views import parse_views  # noqa: E402
from agents.sharepoint_workbook import Refused, read_calendar  # noqa: E402
from tests.sharepoint_helpers import (  # noqa: E402
    VIEWS,
    build_workbook,
    column_letter,
    day_column,
    with_results,
)

_PARSED = parse_views(VIEWS)
PLANNING, TEAM_VIEW = _PARSED["planning"], _PARSED["team"]

WEEK = {
    ("Ann Example", "Presence"): {"2026-01-05": "H", "2026-01-06": "A", "2026-01-07": "I",
                                  "2026-01-08": "½H", "2026-01-09": "T"},
    ("Ann Example", "Guard"): {"2026-01-05": "GDI", "2026-01-06": "GDI", "2026-01-08": "GDI",
                               "2026-01-09": "GDI"},
    ("Bob Sample", "Presence"): {"2026-01-05": "T", "2026-01-06": "X?", "2026-01-07": 3},
    ("Bob Sample", "Guard"): {"2026-01-05": "GDH", "2026-01-06": "GDH", "2026-01-07": "GDH"},
    ("Zed Unlisted", "Guard"): {"2026-01-10": "GDI"},
}


def _read(data, lo="2026-01-05", hi="2026-01-11", view=PLANNING, lookup=TEAM_VIEW):
    return read_calendar(data, view, date.fromisoformat(lo), date.fromisoformat(hi), lookup)


def _runs(out, member, kind):
    return [(r["status"], r["from"], r["to"]) for r in out["runs"]
            if r["member"] == member and r["kind"] == kind]


def test_consecutive_days_with_one_status_are_one_run():
    out = _read(build_workbook({2026: WEEK}))
    assert _runs(out, "Ann Example", "Guard") == [
        ("GDI", "2026-01-05", "2026-01-06"), ("GDI", "2026-01-08", "2026-01-09")]
    assert _runs(out, "Bob Sample", "Guard") == [("GDH", "2026-01-05", "2026-01-07")]
    assert out["sheets"] == ["2026"]


def test_an_absence_reason_is_never_distinguishable():
    out = _read(build_workbook({2026: WEEK}))
    # H, A and I are one run of "unavailable": the reason does not pass.
    assert _runs(out, "Ann Example", "Presence") == [
        ("unavailable", "2026-01-05", "2026-01-07"),
        ("half_day", "2026-01-08", "2026-01-08"),
        ("available", "2026-01-09", "2026-01-09")]
    text = json.dumps(out, ensure_ascii=False)
    assert '"I"' not in text and '"A"' not in text and "½" not in text


def test_a_value_outside_the_code_map_is_counted_and_not_passed_on():
    out = _read(build_workbook({2026: WEEK}))
    assert out["unmapped"] == 2
    assert _runs(out, "Bob Sample", "Presence") == [("available", "2026-01-05", "2026-01-05")]
    assert "X?" not in json.dumps(out)


def test_the_lookup_adds_pinned_columns_and_counts_misses():
    out = _read(build_workbook({2026: WEEK}))
    by_member = {r["member"]: r["ID"] for r in out["runs"]}
    assert by_member == {"Ann Example": 1001, "Bob Sample": 1002, "Zed Unlisted": None}
    assert out["lookup_misses"] == 1


def test_a_name_that_is_in_the_table_twice_is_not_looked_up():
    team = [("Ann Example", "Basis", 1001, "u"), ("Ann Example", "Dev", 1009, "u")]
    out = _read(build_workbook({2026: WEEK}, team=team))
    assert {r["ID"] for r in out["runs"]} == {None}


def test_conflicts_are_guard_days_on_which_the_member_is_not_there():
    out = _read(build_workbook({2026: WEEK}))
    assert out["conflicts"] == [
        {"member": "Ann Example", "from": "2026-01-05", "to": "2026-01-06",
         "status": "GDI", "against": "unavailable"},
        {"member": "Ann Example", "from": "2026-01-08", "to": "2026-01-08",
         "status": "GDI", "against": "half_day"},
    ]


def test_rows_that_cannot_be_used_are_skipped_and_counted():
    extra = {2026: [(None, "Basis", "Guard"), ("#N/A", "Basis", "Presence"),
                    ("Ann Example", "Basis", "Guard")]}
    out = _read(build_workbook({2026: WEEK}, extra_rows=extra))
    assert out["skipped_rows"] == 3
    assert len(_runs(out, "Ann Example", "Guard")) == 2


def test_rows_below_the_stop_marker_are_not_read():
    out = _read(build_workbook({2026: WEEK}), lo="2026-01-01", hi="2026-01-11")
    assert all(r["member"] != "Below the stop marker" for r in out["runs"])
    no_stop = copy.deepcopy(VIEWS)
    del no_stop["planning"]["stop_at"]
    out = _read(build_workbook({2026: WEEK}), lo="2026-01-01", hi="2026-01-11",
                view=parse_views(no_stop)["planning"])
    assert any(r["member"] == "Below the stop marker" for r in out["runs"])


def test_a_window_over_new_year_reads_two_sheets_and_joins_the_run():
    data = build_workbook({
        2026: {("Ann Example", "Guard"): {"2026-12-28": "GDI", "2026-12-29": "GDI",
                                          "2026-12-30": "GDI", "2026-12-31": "GDI"}},
        2027: {("Ann Example", "Guard"): {"2027-01-01": "GDI", "2027-01-02": "GDI",
                                          "2027-01-03": "GDI"}},
    })
    out = _read(data, lo="2026-12-28", hi="2027-01-03")
    assert out["sheets"] == ["2026", "2027"]
    assert _runs(out, "Ann Example", "Guard") == [("GDI", "2026-12-28", "2027-01-03")]


def test_a_year_without_a_sheet_is_refused():
    with pytest.raises(Refused) as refused:
        _read(build_workbook({2026: WEEK}), lo="2026-12-28", hi="2027-01-03")
    assert refused.value.code == "sheet_not_found"


def test_formulas_without_stored_results_are_refused_never_guessed():
    data = build_workbook({2026: WEEK}, date_formulas=True)
    with pytest.raises(Refused) as refused:
        _read(data)
    assert refused.value.code == "no_cached_values"


def test_date_formulas_with_stored_results_are_read():
    data = build_workbook({2026: WEEK}, date_formulas=True)
    # Excel's serial numbers of 5 and 6 January 2026.
    data = with_results(data, "2026", {column_letter(date(2026, 1, 5)) + "8": 46027,
                                       column_letter(date(2026, 1, 6)) + "8": 46028})
    out = _read(data, hi="2026-01-06")
    assert _runs(out, "Ann Example", "Guard") == [("GDI", "2026-01-05", "2026-01-06")]
    # A day whose date formula has no result is not guessed from its column.
    with pytest.raises(Refused) as refused:
        _read(data, hi="2026-01-07")
    assert refused.value.code == "no_cached_values"


def test_a_date_row_that_holds_no_dates_is_layout_drift():
    moved = copy.deepcopy(VIEWS)
    moved["planning"].update(date_row=7, first_row=10)
    with pytest.raises(Refused) as refused:
        _read(build_workbook({2026: WEEK}), view=parse_views(moved)["planning"])
    assert refused.value.code == "dates_not_found"


# --- fix round 1 ------------------------------------------------------------

MARKER = "zz-planted-cell-text"


def _view(*drop, **change):
    views = copy.deepcopy(VIEWS)
    views["planning"].update(change)
    for key in drop:
        del views["planning"][key]
    return parse_views(views)["planning"]


def _refused(*args, **kwargs) -> Refused:
    with pytest.raises(Refused) as refused:
        _read(*args, **kwargs)
    return refused.value


def _members(out):
    return {r["member"] for r in out["runs"]}


# Default-deny on label text (AMENDMENT 1).

def test_a_row_of_a_kind_that_is_not_pinned_is_skipped_and_its_text_goes_nowhere():
    cells = {**WEEK, ("Ann Example", "Sick leave " + MARKER): {"2026-01-05": "H"},
             ("Cy Placeholder", "guard"): {"2026-01-05": "GDI"},
             ("Cy Placeholder", None): {"2026-01-06": "GDI"}}
    out = _read(build_workbook({2026: cells}))
    assert out["skipped_rows"] == 3
    assert {r["kind"] for r in out["runs"]} == {"Presence", "Guard"}
    text = json.dumps(out, ensure_ascii=False)
    assert MARKER not in text and "Sick" not in text and "Cy Placeholder" not in text
    # The rows of the pinned kinds are read as before.
    assert out["runs"] == _read(build_workbook({2026: WEEK}))["runs"]


def test_a_kind_is_compared_trimmed_and_composed():
    cells = {("Ann Example", "  Presence "): {"2026-01-05": "H"},
             ("Bob Sample", "Pre\u0301sence"): {"2026-01-05": "H"}}
    view = _view("conflict", kinds=["Presence", "Pr\u00e9sence"])
    out = _read(build_workbook({2026: cells}), view=view)
    assert {(r["member"], r["kind"]) for r in out["runs"]} == {
        ("Ann Example", "Presence"), ("Bob Sample", "Pr\u00e9sence")}
    assert out["skipped_rows"] == 0


@pytest.mark.parametrize("member", [
    "M" * 121,
    "Ann\tExample " + MARKER,
    "Ann\x7fExample " + MARKER,
    "Ann\x85Example " + MARKER,
    "Ann\nExample " + MARKER,
    "Ann\u2028Example " + MARKER,
    "Ann\u200bExample " + MARKER,   # category Cf
])
def test_a_member_that_is_not_a_bounded_one_line_name_is_skipped(member):
    out = _read(build_workbook({2026: {**WEEK, (member, "Guard"): {"2026-01-05": "GDI"}}}))
    assert out["skipped_rows"] == 1
    text = json.dumps(out, ensure_ascii=False)
    assert MARKER not in text and "MMMM" not in text
    assert _members(out) == {"Ann Example", "Bob Sample", "Zed Unlisted"}


def test_a_member_of_exactly_the_bound_passes():
    member = "M" * 120
    out = _read(build_workbook({2026: {(member, "Guard"): {"2026-01-05": "GDI"}}}))
    assert _members(out) == {member} and out["skipped_rows"] == 0


@pytest.mark.parametrize("team", ["T" * 121, "Ba\u2028sis " + MARKER, "Ba\x7fsis " + MARKER,
                                  "Ba\u2029sis " + MARKER])
def test_a_team_that_is_not_a_bounded_one_line_name_skips_the_row(team):
    data = build_workbook({2026: WEEK}, raw={2026: {"B11": team}})   # Ann, Guard
    out = _read(data)
    assert out["skipped_rows"] == 1
    assert _runs(out, "Ann Example", "Guard") == []
    assert len(_runs(out, "Ann Example", "Presence")) == 3
    text = json.dumps(out, ensure_ascii=False)
    assert MARKER not in text and "TTTT" not in text


def test_a_team_column_that_is_not_configured_is_not_looked_at():
    data = build_workbook({2026: WEEK}, raw={2026: {"B11": "T" * 121}})
    view = _view(labels={"member": "A", "kind": "C"})
    out = _read(data, view=view)
    assert out["skipped_rows"] == 0
    assert {r["team"] for r in out["runs"]} == {None}


def test_a_view_without_a_kind_label_reads_every_row_under_an_empty_kind():
    view = _view("kinds", "conflict", labels={"member": "A", "team": "B"})
    cells = {("Ann Example", "anything " + MARKER): {"2026-01-05": "GDI"}}
    out = _read(build_workbook({2026: cells}), view=view)
    assert [(r["member"], r["kind"], r["status"]) for r in out["runs"]] == [
        ("Ann Example", "", "GDI")]
    assert MARKER not in json.dumps(out)


# The window.

def test_a_reversed_window_is_refused_before_anything_is_opened():
    with pytest.raises(Refused) as refused:
        read_calendar(b"not a workbook", PLANNING, date(2026, 1, 11), date(2026, 1, 5), TEAM_VIEW)
    assert refused.value.code == "invalid_dates"


def test_a_window_that_is_too_long_is_refused_before_anything_is_opened():
    with pytest.raises(Refused) as refused:
        read_calendar(b"not a workbook", PLANNING, date(2026, 1, 1), date(2031, 1, 1), TEAM_VIEW)
    assert refused.value.code == "window_too_long"
    with pytest.raises(Refused) as refused:
        read_calendar(b"not a workbook", PLANNING, date(2026, 1, 1), date(2026, 5, 1), TEAM_VIEW)
    assert refused.value.code == "window_too_long"
    # 120 days is the bound itself: the read gets as far as the file.
    with pytest.raises(Refused) as refused:
        read_calendar(b"not a workbook", PLANNING, date(2026, 1, 1), date(2026, 4, 30), TEAM_VIEW)
    assert refused.value.code == "not_a_workbook"


@pytest.mark.parametrize("lo, hi", [
    ("2026-01-05", date(2026, 1, 11)), (date(2026, 1, 5), None),
    (datetime(2026, 1, 5, 8, 0), datetime(2026, 1, 11, 8, 0)), (20260105, 20260111),
])
def test_a_window_that_is_not_two_dates_is_refused(lo, hi):
    with pytest.raises(Refused) as refused:
        read_calendar(b"not a workbook", PLANNING, lo, hi, TEAM_VIEW)
    assert refused.value.code == "invalid_dates"


def test_leap_day_is_a_day_like_any_other():
    cells = {("Ann Example", "Guard"): {"2028-02-28": "GDI", "2028-02-29": "GDI",
                                        "2028-03-01": "GDI"},
             ("Bob Sample", "Guard"): {"2028-02-29": "GDH"}}
    out = _read(build_workbook({2028: cells}), lo="2028-02-28", hi="2028-03-01")
    assert _runs(out, "Ann Example", "Guard") == [("GDI", "2028-02-28", "2028-03-01")]
    assert _runs(out, "Bob Sample", "Guard") == [("GDH", "2028-02-29", "2028-02-29")]


@pytest.mark.parametrize("day", ["2026-01-01", "2026-12-31"])
def test_a_single_day_at_either_end_of_the_date_row(day):
    cells = {("Ann Example", "Guard"): {"2026-01-01": "GDI", "2026-01-02": "GDH",
                                        "2026-12-30": "GDH", "2026-12-31": "GDI"}}
    out = _read(build_workbook({2026: cells}), lo=day, hi=day)
    assert _runs(out, "Ann Example", "Guard") == [("GDI", day, day)]


# The row cap.

def test_more_rows_than_the_cap_is_a_refusal_not_a_cut(monkeypatch):
    data = build_workbook({2026: WEEK})          # five rows, then the stop marker
    monkeypatch.setattr(sharepoint_workbook, "MAX_CALENDAR_ROWS", 5)
    assert _read(data)["runs"]               # exactly the cap: read
    monkeypatch.setattr(sharepoint_workbook, "MAX_CALENDAR_ROWS", 4)
    assert _refused(data).code == "too_large"
    # Without a stop marker the rows below it count as well.
    monkeypatch.setattr(sharepoint_workbook, "MAX_CALENDAR_ROWS", 6)
    assert _refused(data, lo="2026-01-01", view=_view("stop_at")).code == "too_large"


# Formulas without stored results, per area.

def test_a_body_of_formulas_without_results_is_not_masked_by_a_date_row_with_results():
    cells = {("Ann Example", "Guard"): {"2026-01-05": '=IF(1=1,"GDI","")',
                                        "2026-01-06": '=IF(1=1,"GDI","")'}}
    data = build_workbook({2026: cells}, date_formulas=True)
    data = with_results(data, "2026", {column_letter(date(2026, 1, 5)) + "8": 46027,
                                       column_letter(date(2026, 1, 6)) + "8": 46028})
    assert _refused(data, hi="2026-01-06").code == "no_cached_values"


def test_a_date_row_of_formulas_without_results_is_not_masked_by_the_body():
    cells = {("Ann Example", "Guard"): {"2026-01-05": "=1+1"}}
    data = build_workbook({2026: cells}, date_formulas=True)
    data = with_results(data, "2026", {column_letter(date(2026, 1, 5)) + "10": 2})
    assert _refused(data).code == "no_cached_values"


# The lookup.

def test_a_configured_lookup_that_is_not_handed_in_is_a_programming_error():
    data = build_workbook({2026: WEEK})
    with pytest.raises(ValueError, match="lookup"):
        _read(data, lookup=None)
    with pytest.raises(ValueError, match="lookup"):
        _read(data, lookup=parse_views({"other": VIEWS["team"]})["other"])
    # A view without a lookup needs none.
    out = _read(data, view=_view("lookup"), lookup=None)
    assert out["lookup_misses"] == 0 and all("ID" not in r for r in out["runs"])


# One sheet for every year.

def test_a_sheet_name_without_the_year_is_read_once_over_new_year():
    cells = {("Ann Example", "Guard"): {"2026-12-30": "GDI", "2026-12-31": "GDI",
                                        "2026-12-29": "??"},
             ("", "Guard"): {}}
    # The sheet runs on into January: two more day columns behind 31 December.
    end = day_column(date(2026, 12, 31))
    jan1, jan2 = get_column_letter(end + 1), get_column_letter(end + 2)
    data = build_workbook({2026: cells}, raw={2026: {
        jan1 + "8": datetime(2027, 1, 1), jan2 + "8": datetime(2027, 1, 2),
        jan1 + "10": "GDI", jan2 + "10": "?"}})
    out = _read(data, lo="2026-12-29", hi="2027-01-02", view=_view(sheet="2026"))
    assert out["sheets"] == ["2026"]
    assert out["skipped_rows"] == 1 and out["unmapped"] == 2
    assert _runs(out, "Ann Example", "Guard") == [("GDI", "2026-12-30", "2027-01-01")]


# The date row.

def test_a_requested_day_that_is_in_the_date_row_twice_is_refused():
    data = build_workbook({2026: WEEK}, raw={2026: {
        column_letter(date(2026, 1, 20)) + "8": datetime(2026, 1, 6)}})
    refused = _refused(data)
    assert refused.code == "dates_not_found"
    # A day outside the window may be there twice: it is not read.
    data = build_workbook({2026: WEEK}, raw={2026: {
        column_letter(date(2026, 1, 20)) + "8": datetime(2026, 1, 21)}})
    assert _read(data)["runs"] == _read(build_workbook({2026: WEEK}))["runs"]


def test_a_gap_or_text_in_the_date_row():
    col = column_letter(date(2026, 1, 7))
    for value in (None, "7 jan " + MARKER, "2026-01-07", True):
        data = build_workbook({2026: WEEK}, raw={2026: {col + "8": value}})
        refused = _refused(data)
        assert refused.code == "dates_not_found", value
        # The days around the gap are read.
        assert _runs(_read(data, hi="2026-01-06"), "Ann Example", "Guard") == [
            ("GDI", "2026-01-05", "2026-01-06")]


def test_a_formula_without_a_result_behind_the_dates_does_not_hide_layout_drift():
    # The date row is formulas WITH results for another year's days and one
    # trailing formula without a result (an empty text, a 366th day).
    data = build_workbook({2026: WEEK}, date_formulas=True)
    days = {column_letter(date(2026, 1, 1 + i)) + "8": 45292 + i for i in range(20)}  # 2024
    # All 365 formulas get a result but the last.
    full = {column_letter(date.fromordinal(date(2026, 1, 1).toordinal() + i)) + "8": 45292 + i
            for i in range(364)}
    assert days.items() <= full.items()
    data = with_results(data, "2026", full)
    assert _refused(data).code == "dates_not_found"


# Cells.

def test_what_a_code_cell_holds_never_passes_as_typed():
    cells = {("Ann Example", "Guard"): {
        "2026-01-05": "#N/A", "2026-01-06": True, "2026-01-07": datetime(2026, 3, 4, 5, 6),
        "2026-01-08": "   ", "2026-01-09": 4.5, "2026-01-10": "GDI", "2026-01-11": " GDI "}}
    out = _read(build_workbook({2026: cells}))
    # Error value, boolean, datetime and number are counted; whitespace is an empty cell.
    assert out["unmapped"] == 4
    assert _runs(out, "Ann Example", "Guard") == [("GDI", "2026-01-10", "2026-01-11")]
    text = json.dumps(out)
    assert "N/A" not in text and "true" not in text.lower() and "03-04" not in text \
        and "4.5" not in text


def test_a_conflict_over_days_with_different_codes_of_one_status_is_one_span():
    cells = {("Ann Example", "Presence"): {"2026-01-05": "H", "2026-01-06": "A",
                                           "2026-01-07": "I"},
             ("Ann Example", "Guard"): {"2026-01-05": "GDI", "2026-01-06": "GDI",
                                        "2026-01-07": "GDI", "2026-01-08": "GDI"}}
    out = _read(build_workbook({2026: cells}))
    assert out["conflicts"] == [{"member": "Ann Example", "from": "2026-01-05",
                                 "to": "2026-01-07", "status": "GDI",
                                 "against": "unavailable"}]


def test_merged_and_empty_label_cells_are_rows_without_a_label():
    # Ann's name is one merged cell over her two rows (Presence, Guard); the
    # team cell of Bob's first row is empty.
    data = build_workbook({2026: WEEK}, raw={2026: {"B12": None}},
                          merges={2026: ["A10:A11"]})
    out = _read(data)
    # The second row of the merge has no member of its own: skipped, not guessed.
    assert out["skipped_rows"] == 1
    assert _runs(out, "Ann Example", "Guard") == []
    assert len(_runs(out, "Ann Example", "Presence")) == 3
    assert {r["team"] for r in out["runs"] if r["member"] == "Bob Sample"
            and r["kind"] == "Presence"} == {None}
    # A merged kind cell: the second row has no kind, which is not a pinned one.
    data = build_workbook({2026: WEEK}, merges={2026: ["C12:C13"]})
    out = _read(data)
    assert out["skipped_rows"] == 1 and _runs(out, "Bob Sample", "Guard") == []


# A refusal carries no cell content.

def _marked(**kwargs):
    cells = {("Ann " + MARKER, "Guard"): {"2026-01-05": MARKER, "2026-01-06": "GDI"},
             ("Bob Sample", MARKER): {"2026-01-05": "GDI"}}
    raw = {"B10": "Team " + MARKER, "D7": MARKER, "C8": MARKER, **kwargs.pop("raw", {})}
    return build_workbook({2026: cells}, raw={2026: raw}, **kwargs)


def _refusals(monkeypatch):
    col = column_letter(date(2026, 1, 7))
    yield "sheet_not_found", _refused(_marked(), lo="2027-01-05", hi="2027-01-06")
    yield "sheet_not_found", _refused(_marked(), view=_view(sheet="Plan"))
    yield "dates_not_found", _refused(_marked(raw={col + "8": MARKER}))
    yield "dates_not_found", _refused(_marked(raw={col + "8": datetime(2026, 1, 6)}))
    yield "dates_not_found", _refused(_marked(), view=_view(date_row=7))
    yield "no_cached_values", _refused(_marked(date_formulas=True))
    yield "no_cached_values", _refused(_marked(raw={"E10": '="%s"' % MARKER}))
    yield "invalid_dates", _refused(_marked(), lo="2026-01-11", hi="2026-01-05")
    yield "window_too_long", _refused(_marked(), lo="2026-01-01", hi="2026-12-31")
    monkeypatch.setattr(sharepoint_workbook, "MAX_CALENDAR_ROWS", 1)
    yield "too_large", _refused(_marked())


def test_a_refusal_never_carries_what_a_cell_holds(monkeypatch):
    seen = set()
    for code, refused in _refusals(monkeypatch):
        assert refused.code == code
        seen.add(code)
        said = json.dumps([refused.code, refused.message, refused.hint, str(refused),
                           refused.as_error()])
        assert MARKER not in said and "Ann" not in said and "Bob" not in said, code
    assert seen == {"sheet_not_found", "dates_not_found", "no_cached_values",
                    "invalid_dates", "window_too_long", "too_large"}
