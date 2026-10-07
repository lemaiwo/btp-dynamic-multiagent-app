"""A planning workbook generated for the ``builtin:sharepoint`` tests.

Synthetic names only. The layout is the kind the views describe: a ``Team``
sheet with the Excel table ``TeamMembers``, and one sheet per year with the
dates in row 8 from column D and one row per member and kind from row 10,
ended by a ``Summary`` row.
"""

from __future__ import annotations

import io
import re
import zipfile
from datetime import date, datetime, timedelta
from typing import Any

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table

TEAM = [
    ("Ann Example", "Basis", 1001, "u-ann"),
    ("Bob Sample", "Basis", 1002, "u-bob"),
    ("Cy Placeholder", "Dev", 1003, "u-cy"),
]

VIEWS: dict[str, Any] = {
    "team": {"kind": "table", "table": "TeamMembers", "columns": ["Name", "Team", "ID"]},
    "planning": {
        "kind": "calendar", "sheet": "{year}", "date_row": 8, "first_row": 10,
        "first_date_column": "D",
        "labels": {"member": "A", "team": "B", "kind": "C"},
        "stop_at": "Summary",
        "codes": {"H": "unavailable", "A": "unavailable", "I": "unavailable",
                  "½H": "half_day", "T": "available", "GDI": "GDI", "GDH": "GDH"},
        "lookup": {"view": "team", "on": "Name", "add": ["ID"]},
        "conflict": {"kind": "Guard", "against": "Presence",
                     "when": ["unavailable", "half_day"]},
    },
}


def day_column(day: date) -> int:
    """The column of ``day`` in its year sheet (D = 1 January)."""
    return 4 + (day - date(day.year, 1, 1)).days


def build_workbook(
    cells: dict[int, dict[tuple[str, str], dict[str, Any]]] | None = None,
    *,
    team: list[tuple[Any, ...]] | None = None,
    date_formulas: bool = False,
    extra_rows: dict[int, list[tuple[Any, ...]]] | None = None,
) -> bytes:
    """Workbook bytes. ``cells[year][(member, kind)] = {"2026-01-05": "H"}``.

    ``date_formulas`` writes the date row as formulas, which openpyxl stores
    without a result (see :func:`with_results`). ``extra_rows[year]`` are raw
    ``(A, B, C)`` label rows written after the members, before ``Summary``.
    """
    book = Workbook()
    sheet = book.active
    sheet.title = "Team"
    sheet.append(["Team planning"])
    sheet.append([])
    sheet.append(["Name", "Team", "ID", "UserID"])
    rows = TEAM if team is None else team
    for row in rows:
        sheet.append(list(row))
    sheet.add_table(Table(displayName="TeamMembers", ref=f"A3:D{3 + len(rows)}"))
    for year, lines in (cells or {}).items():
        ws = book.create_sheet(str(year))
        days = (date(year + 1, 1, 1) - date(year, 1, 1)).days
        for i in range(days):
            col = 4 + i
            if date_formulas:
                ws.cell(8, col, "=DATE(%d,1,1)+%d" % (year, i))
            else:
                ws.cell(8, col, datetime(year, 1, 1) + timedelta(days=i))
        r = 10
        for (member, kind), values in lines.items():
            ws.cell(r, 1, member)
            ws.cell(r, 2, "Basis")
            ws.cell(r, 3, kind)
            for iso, value in values.items():
                ws.cell(r, day_column(date.fromisoformat(iso)), value)
            r += 1
        for labels in (extra_rows or {}).get(year, []):
            for c, value in enumerate(labels, 1):
                ws.cell(r, c, value)
            r += 1
        ws.cell(r, 1, "Summary")
        ws.cell(r + 1, 1, "Below the stop marker")
        ws.cell(r + 1, 3, "Presence")
        ws.cell(r + 1, 4, "H")
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def with_results(data: bytes, sheet_title: str, results: dict[str, Any]) -> bytes:
    """``data`` with a stored result for the formula cells in ``results``
    (``{"D8": 46023}``), as an application that calculates would save it."""
    book_xml = zipfile.ZipFile(io.BytesIO(data)).read("xl/workbook.xml").decode()
    names = re.findall(r'<sheet [^>]*name="([^"]+)"', book_xml)
    part = f"xl/worksheets/sheet{names.index(sheet_title) + 1}.xml"
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as src, \
            zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            body = src.read(info.filename)
            if info.filename == part:
                text = body.decode()
                for ref, value in results.items():
                    text, n = re.subn(
                        r'(<c r="%s"[^>]*>\s*<f>[^<]*</f>)\s*<v\s*/?>(</v>)?' % ref,
                        r"\g<1><v>%s</v>" % value, text)
                    assert n == 1, ref
                body = text.encode()
            dst.writestr(info, body)
    return out.getvalue()


def column_letter(day: date) -> str:
    return get_column_letter(day_column(day))
