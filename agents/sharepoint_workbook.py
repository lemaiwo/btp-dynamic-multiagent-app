"""Reads the views of a ``builtin:sharepoint`` entry out of workbook bytes.

Pure: bytes in, plain data out, no HTTP and no config parsing (the views come
validated from :mod:`agents.sharepoint_views`). It is CPU work on a file
somebody else wrote: call it through ``asyncio.to_thread``.

What the file may do to us is bounded before and while it is read: the size,
the number of archive members and their declared uncompressed size, in total
and per part that is read whole (:func:`check_archive`), the rows and columns
of a view, and ``openpyxl`` in read-only mode (the rows of a sheet are
streamed; with ``defusedxml`` installed, which ``requirements.txt`` pins, and
without ``lxml``, a part with an entity declaration is refused, in whatever
part it stands).

Rules of what leaves this module:

* **Table view**: the pinned columns of the Excel table, nothing else, and
  text only as bounded one-line text (:data:`MAX_CELL_CHARS`, no control,
  format or line-separator character): a row with other text in a pinned
  column is skipped and counted.
* **Calendar view, default-deny on cell values**: a cell reaches the caller
  only as the status its code maps to. A value that is not in ``codes`` (an
  unknown code, a number, an error value) is counted as ``unmapped`` and
  never passed on.
* **Calendar view, default-deny on label text**: the kind of a row passes
  only when the view pins it (``kinds``); a member or team passes only as a
  bounded one-line name (:data:`MAX_LABEL_CHARS`, no control, format or
  line-separator character). A sheet is text somebody typed: without these
  two rules a label cell would be a way to put any text before the model.
* **Rows that cannot be used** (no member, an error value in a label cell, a
  kind that is not pinned, a label that is not such a name, a second row for
  the same member and kind) are skipped and counted; their text goes nowhere.
* **No cached values is a refusal.** Formula results are read as stored by
  the application that saved the file. ``openpyxl`` cannot tell a formula
  without a stored result from one whose result is empty, so the rule is:
  a requested day that is missing from the date row while that row holds a
  formula without a result, or a view in which no formula at all has a
  result, is ``no_cached_values``. Nothing is ever calculated or guessed.

Every refusal is a :class:`Refused` with a stable code and a fixed text.
"""

from __future__ import annotations

import io
import posixpath
import re
import unicodedata
import zipfile
from datetime import date, datetime, timedelta
from typing import Any, Iterator

from agents.sharepoint_views import CalendarView, TableView, column_index

__all__ = [
    "MAX_FILE_BYTES", "MAX_LABEL_CHARS", "MAX_MEMBER_BYTES", "MAX_UNCOMPRESSED_BYTES",
    "MAX_WINDOW_DAYS", "Refused", "check_archive", "check_window", "read_calendar",
    "read_table",
]

MAX_FILE_BYTES = 20 * 1024 * 1024
# Declared size of all archive members together, checked before anything is
# inflated. A planning workbook of 3 MB and some thirty sheets inflates to
# about 25 MB; the app container has 1 GB for everything.
MAX_UNCOMPRESSED_BYTES = 64 * 1024 * 1024
# Declared size of one member that is not a sheet. Only the rows of a sheet
# are streamed: the workbook part, the styles, the relationships and a table
# are parsed into a tree several times their size, and the shared strings are
# all kept, on every open of the bytes.
MAX_MEMBER_BYTES = 16 * 1024 * 1024
_SHEETS = "xl/worksheets/"
# One piece of a member, inflated.
_READ_BYTES = 64 * 1024
MAX_ARCHIVE_MEMBERS = 5000
MAX_WINDOW_DAYS = 120
MAX_TABLE_ROWS = 5000
MAX_TABLE_COLUMNS = 200
# Rows of a calendar from ``first_row`` on. Reading ends at ``stop_at`` or
# after ``BLANK_ROWS_END`` rows in a row with nothing in them; a sheet that
# has neither within the cap is refused, never cut off.
MAX_CALENDAR_ROWS = 2000
BLANK_ROWS_END = 50
# A year sheet has at most 366 day columns; a few spare for layout.
MAX_DAY_COLUMNS = 400
# A member or team cell that is longer is not a name: the row is skipped.
MAX_LABEL_CHARS = 120
# A text cell of a pinned table column that is longer skips its row.
MAX_CELL_CHARS = 255

_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_PKG_REL = "http://schemas.openxmlformats.org/package/2006/relationships"
_DOC_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_ERROR = object()  # a cell holding an Excel error value
# ``date.fromisoformat`` alone also reads a week date (``2026-W01-1``).
_DATE_RE = re.compile(r"[0-9]{4}-[0-9]{2}-[0-9]{2}")
_YEAR = "{year}"


class Refused(Exception):
    """A read that is refused: ``code`` is stable, ``message`` a fixed text."""

    def __init__(self, code: str, message: str, hint: str | None = None) -> None:
        super().__init__(message)
        self.code = code
        self.message = message
        self.hint = hint

    def as_error(self) -> dict[str, Any]:
        error: dict[str, Any] = {"code": self.code, "message": self.message}
        if self.hint:
            error["hint"] = self.hint
        return {"error": error}


def check_window(date_from: Any, date_to: Any) -> tuple[date, date]:
    """Two dates as ``YYYY-MM-DD``, in order, at most :data:`MAX_WINDOW_DAYS` days."""
    try:
        if not (isinstance(date_from, str) and isinstance(date_to, str)
                and _DATE_RE.fullmatch(date_from) and _DATE_RE.fullmatch(date_to)):
            raise ValueError
        lo, hi = date.fromisoformat(date_from), date.fromisoformat(date_to)
    except ValueError:
        raise Refused(
            "invalid_dates", "date_from and date_to must be dates as YYYY-MM-DD"
        ) from None
    if hi < lo:
        raise Refused("invalid_dates", "date_to must not be before date_from")
    if (hi - lo).days + 1 > MAX_WINDOW_DAYS:
        raise Refused(
            "window_too_long", f"a window may span at most {MAX_WINDOW_DAYS} days",
            "ask for a shorter period",
        )
    return lo, hi


def check_archive(data: bytes) -> None:
    """Size and archive caps, before the workbook is parsed.

    From the sizes the archive declares: nothing is inflated here.
    :class:`_Archive` holds every later read to them.
    """
    if len(data) > MAX_FILE_BYTES:
        raise Refused("too_large", "the workbook is larger than this tool reads")
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            infos = zf.infolist()
    except (zipfile.BadZipFile, ValueError, OSError):
        raise Refused("not_a_workbook", "the file is not an Excel workbook") from None
    if len(infos) > MAX_ARCHIVE_MEMBERS \
            or sum(i.file_size for i in infos) > MAX_UNCOMPRESSED_BYTES:
        raise Refused("too_large", "the workbook is larger than this tool reads")
    for info in infos:
        # A relationships part below the sheets' folder is read whole too.
        streamed = info.filename.startswith(_SHEETS) and not info.filename.endswith(".rels")
        if not streamed and info.file_size > MAX_MEMBER_BYTES:
            raise Refused("too_large", "the workbook is larger than this tool reads")


class _Member:
    """An archive member as a stream that keeps to its declared size.

    ``ZipExtFile.read()`` without a size inflates the whole compressed stream
    and cuts it to the declared size afterwards, so a member that declares
    100 bytes can put hundreds of megabytes in memory. Read in pieces it stops
    at the declared size. Only what the parsers use is offered.
    """

    def __init__(self, raw: Any, size: int) -> None:
        self._raw = raw
        self._size = size

    def read(self, n: int | None = -1) -> bytes:
        if n is not None and n >= 0:
            return self._raw.read(min(n, _READ_BYTES))
        # Read whole: whatever the part is called, it is not a streamed sheet.
        if self._size > MAX_MEMBER_BYTES:
            raise Refused("too_large", "the workbook is larger than this tool reads")
        pieces = []
        while piece := self._raw.read(_READ_BYTES):
            pieces.append(piece)
        return b"".join(pieces)

    def close(self) -> None:
        self._raw.close()

    def __enter__(self) -> "_Member":
        return self

    def __exit__(self, *exc: Any) -> None:
        self.close()


class _Archive(zipfile.ZipFile):
    """The package as openpyxl and this module read it: see :class:`_Member`.

    ``ZipFile.read`` goes through ``open``, so this covers both.
    """

    def open(self, name: Any, mode: str = "r", *args: Any, **kwargs: Any) -> Any:
        raw = super().open(name, mode, *args, **kwargs)
        if mode != "r":
            return raw
        return _Member(raw, (name if isinstance(name, zipfile.ZipInfo)
                             else self.getinfo(name)).file_size)


def _open(data: bytes, *, data_only: bool) -> Any:
    # ``load_workbook`` is these three calls; the reader is built by hand
    # only to hand it the archive that keeps to the declared sizes
    # (``requirements.txt`` pins openpyxl and a test pins the version).
    from openpyxl.reader.excel import ExcelReader

    try:
        reader = ExcelReader(io.BytesIO(data), read_only=True, data_only=data_only,
                             keep_links=False)
        reader.archive.close()
        reader.archive = _Archive(io.BytesIO(data))
        reader.read()
        return reader.wb
    except Refused:
        raise
    except Exception:  # noqa: BLE001 - whatever the parser says, the answer is one code
        raise Refused("not_a_workbook", "the file is not an Excel workbook") from None


def _plain(cell: Any) -> Any:
    """A cell's stored value as plain data; :data:`_ERROR` for an error value."""
    if getattr(cell, "data_type", "n") == "e":
        return _ERROR
    value = cell.value
    if value is None or isinstance(value, bool):
        return value
    if isinstance(value, str):
        # A cell of whitespace only is an empty cell: in a calendar it is
        # neither a code nor counted as ``unmapped``.
        text = unicodedata.normalize("NFC", value).strip()
        return text or None
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == datetime.min.time() \
            else value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, (int, float)):
        return value
    return str(value)


def _rows(
    ws: Any, min_row: int, max_row: int, max_col: int, min_col: int = 1
) -> Iterator[tuple]:
    try:
        yield from ws.iter_rows(min_row=min_row, max_row=max_row,
                                min_col=min_col, max_col=max_col)
    except Refused:
        raise
    except Exception:  # noqa: BLE001 - a broken sheet part
        raise Refused("not_a_workbook", "the file is not an Excel workbook") from None


def _no_results(formulas: int, with_result: int) -> None:
    if formulas and not with_result:
        raise Refused(
            "no_cached_values",
            "the workbook holds formulas without stored results",
            "it was last saved by a program that does not calculate; "
            "open and save it in Excel",
        )


# --- table view -----------------------------------------------------------

def _rels(zf: zipfile.ZipFile, part: str) -> dict[str, tuple[str, str]]:
    """Relationship id -> (type, target part) of ``part`` ('' = the package)."""
    from openpyxl.xml.functions import fromstring

    folder, name = posixpath.split(part)
    try:
        root = fromstring(zf.read(posixpath.join(folder, "_rels", f"{name}.rels")))
    except KeyError:
        return {}
    out: dict[str, tuple[str, str]] = {}
    for rel in root.iter(f"{{{_PKG_REL}}}Relationship"):
        if rel.get("TargetMode") == "External":
            continue
        target = rel.get("Target") or ""
        path = target.lstrip("/") if target.startswith("/") \
            else posixpath.normpath(posixpath.join(folder, target))
        out[rel.get("Id") or ""] = (rel.get("Type") or "", path)
    return out


def _find_table(data: bytes, name: str) -> tuple[str, str, int]:
    """(sheet title, cell range, totals rows) of the Excel table ``name``.

    Read from the package parts: a read-only worksheet of openpyxl does not
    load its tables.
    """
    from openpyxl.xml.functions import fromstring

    try:
        with _Archive(io.BytesIO(data)) as zf:
            book = next((p for t, p in _rels(zf, "").values()
                         if t.endswith("/officeDocument")), "xl/workbook.xml")
            parts = _rels(zf, book)
            for sheet in fromstring(zf.read(book)).iter(f"{{{_MAIN}}}sheet"):
                part = parts.get(sheet.get(f"{{{_DOC_REL}}}id") or "")
                if part is None:
                    continue
                for kind, target in _rels(zf, part[1]).values():
                    if not kind.endswith("/table"):
                        continue
                    table = fromstring(zf.read(target))
                    if name in (table.get("name"), table.get("displayName")):
                        return (sheet.get("name") or "", table.get("ref") or "",
                                int(table.get("totalsRowCount") or 0))
    except Refused:
        raise
    except Exception:  # noqa: BLE001 - a broken package
        raise Refused("not_a_workbook", "the file is not an Excel workbook") from None
    raise Refused("table_not_found", "the workbook has no table of the configured name",
                  "an administrator must correct the view")


def read_table(data: bytes, view: TableView) -> dict[str, Any]:
    """The rows of a table view: the pinned columns only."""
    from openpyxl.utils.cell import range_boundaries

    check_archive(data)
    title, ref, totals = _find_table(data, view.table)
    try:
        min_col, min_row, max_col, max_row = range_boundaries(ref)
    except Exception:  # noqa: BLE001
        raise Refused("not_a_workbook", "the file is not an Excel workbook") from None
    max_row -= totals
    if max_row - min_row > MAX_TABLE_ROWS or max_col - min_col >= MAX_TABLE_COLUMNS:
        raise Refused("too_large", "the table is larger than this tool reads")

    book = _open(data, data_only=True)
    try:
        if title not in book.sheetnames:
            raise Refused("table_not_found", "the workbook has no table of the configured name")
        grid = [[_plain(c) for c in row]
                for row in _rows(book[title], min_row, max_row, max_col, min_col)]
    finally:
        book.close()
    header = grid[0] if grid else []
    index: dict[str, int] = {}
    for i, name in enumerate(header):
        if isinstance(name, str):
            index.setdefault(name, i)
    if any(c not in index for c in view.columns):
        raise Refused("column_not_found", "a pinned column is not in the table",
                      "an administrator must correct the view")

    formulas = with_result = 0
    book = _open(data, data_only=False)
    try:
        for r, row in enumerate(_rows(book[title], min_row + 1, max_row, max_col, min_col), 1):
            for c, cell in enumerate(row):
                if cell.data_type == "f":
                    formulas += 1
                    with_result += r < len(grid) and c < len(grid[r]) and grid[r][c] is not None
    finally:
        book.close()
    _no_results(formulas, with_result)

    rows = []
    skipped = 0
    for line in grid[1:]:
        record = {c: (None if line[index[c]] is _ERROR else line[index[c]])
                  for c in view.columns}
        # Text is what somebody typed: it passes only as bounded one-line
        # text, and a row that holds other text is not returned in part.
        if any(isinstance(v, str) and not _is_name(v, MAX_CELL_CHARS)
               for v in record.values()):
            skipped += 1
        elif any(v is not None for v in record.values()):
            rows.append(record)
    return {"columns": list(view.columns), "rows": rows, "skipped_rows": skipped}


# --- calendar view --------------------------------------------------------

def _as_date(cell: Any, epoch: Any) -> date | None:
    value = cell.value
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if type(value) in (int, float) and 1 <= value < 2_958_466:
        from openpyxl.utils.datetime import from_excel

        try:
            moment = from_excel(value, epoch)
        except Exception:  # noqa: BLE001
            return None
        return moment.date() if isinstance(moment, datetime) else None
    return None


def _is_name(text: str, limit: int = MAX_LABEL_CHARS) -> bool:
    """Whether cell text may pass (a member or team; with ``limit``, a table
    cell): bounded, one line, no control or format character (categories C*,
    Zl, Zp)."""
    return len(text) <= limit and not any(
        unicodedata.category(ch).startswith("C")
        or unicodedata.category(ch) in ("Zl", "Zp") for ch in text)


def _dates_not_found() -> Refused:
    return Refused(
        "dates_not_found",
        "the date row of the sheet does not hold every requested day exactly once",
        "the layout of the sheet may differ from the configured view",
    )


def _read_sheet(
    data: bytes, view: CalendarView, title: str, lo: date, hi: date, state: dict[str, Any]
) -> None:
    """Adds the days ``lo``..``hi`` of the sheet ``title`` to ``state``."""
    first_col = column_index(view.first_date_column)
    last_col = first_col + MAX_DAY_COLUMNS - 1
    label_cols = {k: column_index(c) for k, c in view.labels.items()}
    wanted = {lo + timedelta(days=i) for i in range((hi - lo).days + 1)}

    book = _open(data, data_only=True)
    try:
        if title not in book.sheetnames:
            if _YEAR in view.sheet:
                raise Refused("sheet_not_found", f"the workbook has no sheet for {lo.year}",
                              "ask for a period the workbook covers")
            raise Refused("sheet_not_found", "the workbook has no sheet of the configured name",
                          "an administrator must correct the view")
        ws = book[title]
        day_cols: dict[date, int] = {}
        twice = False
        date_values: dict[int, bool] = {}
        last_dated = 0
        for row in _rows(ws, view.date_row, view.date_row, last_col, first_col):
            for offset, cell in enumerate(row):
                day = _as_date(cell, book.epoch)
                date_values[first_col + offset] = cell.value is not None
                if day is not None:
                    last_dated = first_col + offset
                if day in wanted:
                    # The same day in two columns: which one is meant is a guess.
                    twice = twice or day in day_cols
                    day_cols.setdefault(day, first_col + offset)
        max_col = max(day_cols.values(), default=first_col)
        lines: list[tuple[int, list[Any]]] = []
        blank = 0
        stop = view.stop_at.casefold() if view.stop_at else None
        # Up to BLANK_ROWS_END rows past the cap, so that a sheet which ends
        # (stop marker, blank rows) right at the cap is not taken for one that
        # goes on.
        for n, row in enumerate(_rows(
                ws, view.first_row,
                view.first_row + MAX_CALENDAR_ROWS + BLANK_ROWS_END - 1, max_col)):
            values = [_plain(c) for c in row]
            values += [None] * (max_col - len(values))
            labels = [values[c - 1] for c in label_cols.values()]
            if stop and any(isinstance(v, str) and v.casefold() == stop for v in labels):
                break
            if all(v is None for v in values):
                blank += 1
                if blank >= BLANK_ROWS_END:
                    break
                continue
            if n >= MAX_CALENDAR_ROWS:
                raise Refused("too_large", "the sheet has more rows than this tool reads",
                              "an administrator can end the rows with stop_at")
            blank = 0
            lines.append((view.first_row + n, values))
    finally:
        book.close()

    # Second pass, formulas as written: which cells are formulas at all. The
    # date row and the rows below it are two areas, each decided on its own:
    # a date row with results must not hide a body that has none.
    date_formulas = date_results = body_formulas = body_results = 0
    date_formula_without_result = False
    by_row = dict(lines)
    book = _open(data, data_only=False)
    try:
        ws = book[title]
        for row in _rows(ws, view.date_row, view.date_row, last_col, first_col):
            for offset, cell in enumerate(row):
                if cell.data_type == "f":
                    date_formulas += 1
                    if date_values.get(first_col + offset):
                        date_results += 1
                    elif first_col + offset < last_dated:
                        # Only a hole among the dates. A formula behind the
                        # last date (an empty 366th day) says nothing about
                        # the days that are missing.
                        date_formula_without_result = True
        if lines:
            for n, row in enumerate(_rows(ws, view.first_row, lines[-1][0], max_col)):
                values = by_row.get(view.first_row + n)
                for c, cell in enumerate(row):
                    if cell.data_type == "f":
                        body_formulas += 1
                        body_results += values is not None and values[c] is not None
    finally:
        book.close()
    _no_results(date_formulas, date_results)
    _no_results(body_formulas, body_results)
    if wanted - set(day_cols):
        if date_formula_without_result:
            _no_results(1, 0)
        raise _dates_not_found()
    if twice:
        raise _dates_not_found()

    has_kind, has_team = "kind" in label_cols, "team" in label_cols
    seen: set[tuple[str, str]] = set()
    for _, values in lines:
        label = {k: values[c - 1] for k, c in label_cols.items()}
        member, kind, team = label.get("member"), label.get("kind"), label.get("team")
        usable = isinstance(member, str) and _is_name(member) \
            and not any(v is _ERROR for v in label.values())
        # The kind of a row is exactly one of the pinned ones, or the row is
        # not read (``_plain`` has trimmed and composed the cell).
        if usable and has_kind:
            usable = isinstance(kind, str) and kind in view.kinds
        if usable and has_team and isinstance(team, str):
            usable = _is_name(team)
        key = (member, kind if has_kind else "")
        if not usable or key in seen:
            state["skipped_rows"] += 1
            continue
        seen.add(key)
        row_state = state["rows"].setdefault(
            key, {"member": member, "kind": key[1],
                  "team": team if isinstance(team, str) else None, "days": {}})
        for day, col in day_cols.items():
            value = values[col - 1]
            if value is None:
                continue
            status = None if value is _ERROR or isinstance(value, bool) \
                else view.codes.get(str(value))
            if status is None:
                state["unmapped"] += 1
            else:
                row_state["days"][day] = status


def _runs(days: dict[date, Any]) -> Iterator[tuple[date, date, Any]]:
    """Consecutive days with the same value, as (from, to, value)."""
    start = prev = None
    for day in sorted(days):
        if start is not None and day - prev == timedelta(days=1) and days[day] == days[start]:
            prev = day
            continue
        if start is not None:
            yield start, prev, days[start]
        start = prev = day
    if start is not None:
        yield start, prev, days[start]


def read_calendar(
    data: bytes,
    view: CalendarView,
    date_from: date,
    date_to: date,
    lookup: TableView | None = None,
) -> dict[str, Any]:
    """The runs of a calendar view between two dates (inclusive).

    ``lookup`` is the table view ``view.lookup`` names; the caller resolves
    it from the same config. A view with a lookup read without that table
    view is a mistake of the caller and raises ``ValueError`` (not a
    :class:`Refused`, which is about the file or the request): the runs would
    silently lack the pinned columns.

    The window is checked here as well as by the caller
    (:func:`check_window`): reversed or longer than
    :data:`MAX_WINDOW_DAYS` is refused before anything is opened.
    """
    if type(date_from) is not date or type(date_to) is not date:
        raise Refused("invalid_dates", "date_from and date_to must be dates as YYYY-MM-DD")
    date_from, date_to = check_window(date_from.isoformat(), date_to.isoformat())
    if view.lookup is not None and (lookup is None or lookup.name != view.lookup.view):
        raise ValueError("read_calendar: the view has a lookup; pass its table view")
    check_archive(data)
    state: dict[str, Any] = {"rows": {}, "skipped_rows": 0, "unmapped": 0}
    # One read per sheet: a sheet name without {year} is the same sheet for
    # every year of the window and is read once, with the whole window.
    spans: dict[str, tuple[date, date]] = {}
    for year in range(date_from.year, date_to.year + 1):
        title = view.sheet.replace(_YEAR, str(year))
        lo = max(date_from, date(year, 1, 1)) if _YEAR in view.sheet else date_from
        hi = min(date_to, date(year, 12, 31)) if _YEAR in view.sheet else date_to
        spans.setdefault(title, (lo, hi))
    for title, (lo, hi) in spans.items():
        _read_sheet(data, view, title, lo, hi, state)
    sheets = list(spans)

    extra: dict[str, dict[str, Any]] = {}
    add: tuple[str, ...] = ()
    if view.lookup is not None:
        add = view.lookup.add
        seen: dict[str, int] = {}
        table = read_table(data, lookup)["rows"]
        for record in table:
            key = record.get(view.lookup.on)
            if isinstance(key, str):
                seen[key] = seen.get(key, 0) + 1
                extra[key] = record
        # A name that is in the table twice is not a key: no guess.
        extra = {k: v for k, v in extra.items() if seen[k] == 1}

    runs: list[dict[str, Any]] = []
    misses: set[str] = set()
    for row in state["rows"].values():
        found = extra.get(row["member"])
        if add and found is None and row["days"]:
            misses.add(row["member"])
        for start, end, status in _runs(row["days"]):
            run = {"member": row["member"], "team": row["team"], "kind": row["kind"],
                   "status": status, "from": start.isoformat(), "to": end.isoformat()}
            for column in add:
                run[column] = found.get(column) if found else None
            runs.append(run)
    runs.sort(key=lambda r: (r["member"], r["kind"], r["from"]))

    conflicts: list[dict[str, Any]] = []
    if view.conflict is not None:
        for (member, row_kind), row in state["rows"].items():
            other = state["rows"].get((member, view.conflict.against))
            if row_kind != view.conflict.kind or other is None:
                continue
            clash = {day: (status, other["days"][day])
                     for day, status in row["days"].items()
                     if other["days"].get(day) in view.conflict.when}
            for start, end, (status, other_status) in _runs(clash):
                conflicts.append({"member": member, "from": start.isoformat(),
                                  "to": end.isoformat(), "status": status,
                                  "against": other_status})
        conflicts.sort(key=lambda c: (c["member"], c["from"]))

    return {
        "runs": runs,
        "conflicts": conflicts,
        "skipped_rows": state["skipped_rows"],
        "unmapped": state["unmapped"],
        "lookup_misses": len(misses),
        "sheets": sheets,
    }
